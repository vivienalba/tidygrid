"""Run: python -m unittest test_reporting.py."""

import io, unittest
from datetime import date
import pandas as pd
from openpyxl import load_workbook
from reporting import period_bounds, load_file, build_report
from report_exports import excel_bytes, pdf_bytes


class ReportingTests(unittest.TestCase):
    def setUp(self):
        self.data = pd.DataFrame(
            {
                "date": ["2024-02-01", "2024-02-29", "2024-03-01", "bad", "2024-02-14"],
                "task": ["A", "B", "C", "D", '=HYPERLINK("x")'],
                "status": ["Done", "Pending", "Done", "Done", "Cancelled"],
                "owner": ["Alex", "Sam", "Alex", "Sam", "Alex"],
                "due": ["2024-02-03", "2024-02-20", "", "", "2024-02-15"],
                "completed": ["2024-02-02", "", "2024-03-02", "", ""],
                "value": ["10", "20", "30", "40", "not a number"],
            }
        )
        self.mapping = {
            "Activity date": "date",
            "Status": "status",
            "Activity": "task",
            "Owner": "owner",
            "Due date": "due",
            "Completed date": "completed",
            "Category": None,
            "Value": "value",
        }

    def report(self, anchor=date(2024, 2, 15)):
        return build_report(
            self.data,
            self.mapping,
            "Monthly",
            anchor,
            "YYYY-MM-DD",
            ["Done"],
            ["Cancelled"],
            "Test report",
            "Value total",
        )

    def test_calendar_boundaries(self):
        self.assertEqual(
            period_bounds("Weekly", date(2026, 1, 1)),
            (date(2025, 12, 29), date(2026, 1, 4)),
        )
        self.assertEqual(
            period_bounds("Monthly", date(2024, 2, 1)),
            (date(2024, 2, 1), date(2024, 2, 29)),
        )
        self.assertEqual(
            period_bounds("Quarterly", date(2026, 12, 31)),
            (date(2026, 10, 1), date(2026, 12, 31)),
        )
        self.assertEqual(
            period_bounds("Annual", date(2024, 8, 1)),
            (date(2024, 1, 1), date(2024, 12, 31)),
        )

    def test_counts_and_quality(self):
        r = self.report()
        m = r["metrics"]
        self.assertEqual(m["Activities"], 3)
        self.assertEqual(m["Completed"], 1)
        self.assertEqual(m["Open"], 1)
        self.assertEqual(m["Overdue"], 1)
        self.assertEqual(m["Completion rate (%)"], 50)
        self.assertEqual(m["Value total"], 30)
        self.assertEqual(len(r["issues"]), 1)

    def test_missing_completion_not_invented(self):
        self.data.loc[0, "completed"] = ""
        r = self.report()
        self.assertEqual(r["metrics"]["Completed"], 0)
        self.assertEqual(r["metrics"]["Open"], 2)

    def test_empty_period_export(self):
        r = self.report(date(2023, 2, 15))
        self.assertEqual(r["metrics"]["Activities"], 0)
        self.assertTrue(excel_bytes(r).startswith(b"PK"))
        self.assertTrue(pdf_bytes(r).startswith(b"%PDF"))

    def test_file_formats_preserve_identifiers(self):
        for sep, fn in [(",", "data.csv"), ("\t", "data.tsv")]:
            d = load_file(f"id{sep}name\n001{sep}Viv\n".encode(), fn)
            self.assertEqual(d.iloc[0, 0], "001")
        b = io.BytesIO()
        pd.DataFrame({"id": ["001"], "date": [pd.Timestamp("2024-02-01")]}).to_excel(
            b, index=False
        )
        d = load_file(b.getvalue(), "data.xlsx")
        self.assertEqual(d.iloc[0, 0], "001")

    def test_export_contains_all_rows_and_no_formulas(self):
        r = self.report()
        wb = load_workbook(io.BytesIO(excel_bytes(r)))
        self.assertEqual(wb["Activities"].max_row, len(r["records"]) + 1)
        self.assertTrue(all(c.data_type != "f" for ws in wb for row in ws for c in row))
        self.assertIn("Settings", wb.sheetnames)

    def test_date_formats(self):
        self.data.loc[0, "date"] = "02/01/2024"
        r = build_report(
            self.data,
            self.mapping,
            "Monthly",
            date(2024, 1, 15),
            "DD/MM/YYYY",
            ["Done"],
            ["Cancelled"],
            "Test",
            "Value total",
        )
        self.assertEqual(r["metrics"]["Activities"], 1)


if __name__ == "__main__":
    unittest.main()
