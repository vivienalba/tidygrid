"""Patch only plain-text cell values in an XLSX package; never rebuild its design."""
import io
import posixpath
import re
import zipfile
from xml.etree import ElementTree as ET
from xml.sax.saxutils import escape

import pandas as pd
from openpyxl.utils.cell import range_boundaries, get_column_letter

NS = 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'
REL = 'http://schemas.openxmlformats.org/package/2006/relationships'
RID = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'
CELL = re.compile(rb'<(?P<tag>(?:[\w.-]+:)?c)\b(?P<attrs>[^>]*?)(?:\s*/>|>.*?</(?P=tag)\s*>)', re.S)


def _xml(raw):
    if b'<!DOCTYPE' in raw or b'<!ENTITY' in raw:
        raise ValueError('This workbook contains unsupported XML declarations.')
    return ET.fromstring(raw)


def _target(base, target):
    path = posixpath.normpath(posixpath.join(posixpath.dirname(base), target)) if not target.startswith('/') else target.lstrip('/')
    if path.startswith('../'):
        raise ValueError('Invalid workbook relationship.')
    return path


def _rels(z, part):
    path = posixpath.join(posixpath.dirname(part), '_rels', posixpath.basename(part) + '.rels')
    if path not in z.namelist():
        return {}
    return {r.get('Id'): (_target(part, r.get('Target')), r.get('Type', ''))
            for r in _xml(z.read(path)) if r.get('TargetMode') != 'External'}


class WorkbookSource:
    def __init__(self, raw):
        self.raw = raw
        self.z = zipfile.ZipFile(io.BytesIO(raw))
        names = self.z.namelist()
        if len(names) != len(set(names)):
            raise ValueError('The workbook contains repeated package parts.')
        if any(n.startswith('_xmlsignatures/') for n in names):
            raise ValueError('Digitally signed workbooks cannot be edited without invalidating their signature.')
        if sum(i.file_size for i in self.z.infolist()) > 512 * 1024 * 1024:
            raise ValueError('The expanded workbook is too large to process in this app.')
        root_rels = _xml(self.z.read('_rels/.rels'))
        book = next((_target('', r.get('Target')) for r in root_rels
                     if r.get('Type', '').endswith('/officeDocument')), None)
        if not book:
            raise ValueError('No Excel workbook was found.')
        rels = _rels(self.z, book)
        self.sheets = {}
        for s in _xml(self.z.read(book)).findall(f'{{{NS}}}sheets/{{{NS}}}sheet'):
            relation = rels.get(s.get(f'{{{RID}}}id'))
            if relation and relation[1].endswith('/worksheet'):
                self.sheets[s.get('name')] = relation[0]
        if not self.sheets:
            raise ValueError('No supported worksheets were found. Use a standard .xlsx workbook.')
        shared = next((p for p, t in rels.values() if t.endswith('/sharedStrings')), None)
        self.strings = []
        if shared:
            for si in _xml(self.z.read(shared)):
                text = ''.join(t.text or '' for t in si.iter(f'{{{NS}}}t'))
                rich = si.find(f'{{{NS}}}r') is not None or si.find(f'{{{NS}}}rPh') is not None
                self.strings.append((text, rich))

    def sheet(self, name):
        raw = self.z.read(self.sheets[name])
        root = _xml(raw)
        cells = {}
        for c in root.findall(f'{{{NS}}}sheetData/{{{NS}}}row/{{{NS}}}c'):
            ref = c.get('r')
            if not ref:
                raise ValueError('A cell has no coordinate; this workbook cannot be safely patched.')
            t = c.get('t', 'n')
            v = c.find(f'{{{NS}}}v')
            value, editable = '', False
            if c.find(f'{{{NS}}}f') is not None:
                value = '=' + (c.find(f'{{{NS}}}f').text or '')
            elif t == 's' and v is not None:
                value, rich = self.strings[int(v.text)]
                editable = not rich
            elif t == 'inlineStr':
                inline = c.find(f'{{{NS}}}is')
                if inline is not None:
                    value = ''.join(x.text or '' for x in inline.iter(f'{{{NS}}}t'))
                    editable = inline.find(f'{{{NS}}}r') is None and inline.find(f'{{{NS}}}rPh') is None
            elif v is not None:
                value = v.text or ''
            cells[ref] = (value, editable)
        nonempty = [range_boundaries(k)[:2] for k, (v, _) in cells.items() if v != '']
        if nonempty:
            lo_col, lo_row = min(x[0] for x in nonempty), min(x[1] for x in nonempty)
            hi_col, hi_row = max(x[0] for x in nonempty), max(x[1] for x in nonempty)
            extent = f'{get_column_letter(lo_col)}{lo_row}:{get_column_letter(hi_col)}{hi_row}'
        else:
            extent = 'A1:A1'
        merges = [range_boundaries(m.get('ref')) for m in root.findall(f'{{{NS}}}mergeCells/{{{NS}}}mergeCell')]
        links = [range_boundaries(h.get('ref')) for h in root.findall(f'{{{NS}}}hyperlinks/{{{NS}}}hyperlink')]
        return {'raw': raw, 'cells': cells, 'extent': extent,
                'protected': root.find(f'{{{NS}}}sheetProtection') is not None,
                'skip_ranges': merges + links}


def _patch(raw, changes):
    found = set()
    def replace(match):
        attrs = match.group('attrs')
        ref = re.search(rb'\br\s*=\s*([\x22\x27])([^\x22\x27]+)\1', attrs)
        key = ref.group(2).decode() if ref else None
        if key not in changes:
            return match.group(0)
        found.add(key)
        tag = match.group('tag')
        prefix = tag[:-1]
        attrs = re.sub(rb'\s+t\s*=\s*([\x22\x27])[^\x22\x27]*\1', b'', attrs)
        original = match.group(0)
        inner = original[original.find(b'>') + 1:original.rfind(b'</')] if not original.rstrip().endswith(b'/>') else b''
        p = re.escape(prefix)
        inner = re.sub(rb'<'+p+rb'(?:v|is)\b[^>]*(?:/>|>.*?</'+p+rb'(?:v|is)\s*>)', b'', inner, flags=re.S)
        value = escape(changes[key]).replace('\r', '&#13;').encode('utf-8')
        return b'<'+tag+attrs+b' t="inlineStr"><'+prefix+b'is><'+prefix+b't xml:space="preserve">'+value+b'</'+prefix+b't></'+prefix+b'is>'+inner+b'</'+tag+b'>'
    patched = CELL.sub(replace, raw)
    if found != set(changes):
        raise ValueError('Some cells could not be safely updated; no download was created.')
    return patched


def clean_workbook(source, ranges, options, phone_cleaner):
    patched_parts, reports = {}, {}
    for name, region in ranges.items():
        sheet = source.sheet(name)
        try:
            left, top, right, bottom = range_boundaries(region.upper().replace('$', ''))
            if not all(isinstance(v, int) for v in (left, top, right, bottom)) or left < 1 or top < 1 or left > right or top > bottom or right > 16384 or bottom > 1048576:
                raise ValueError()
            if (bottom-top+1)*(right-left+1) > 1_000_000:
                raise ValueError('Choose a range with at most 1,000,000 cells.')
        except (ValueError, TypeError):
            raise ValueError(f'{name}: enter a valid rectangular range such as A1:D500 (maximum 1,000,000 cells).')
        cells = sheet['cells']
        headers = [str(cells.get(f'{get_column_letter(col)}{top}', ('', False))[0]) or f'Column {get_column_letter(col)}' for col in range(left, right+1)]
        # Preview labels are unique only in the UI; workbook headers remain untouched.
        labels, used = [], set()
        for h in headers:
            label, i = h, 2
            while label in used:
                label = f'{h} ({i})'; i += 1
            labels.append(label); used.add(label)
        changes, log, missing, rows, before_rows, duplicates, seen = {}, [], [], [], [], [], set()
        skipped = 0
        for row in range(top+1, bottom+1):
            before, after = [], []
            for col in range(left, right+1):
                ref = f'{get_column_letter(col)}{row}'
                value, editable = cells.get(ref, ('', False))
                new = value
                blocked = sheet['protected'] or any(a <= col <= c and b <= row <= d for a,b,c,d in sheet['skip_ranges'])
                if editable and not blocked:
                    if options.get('whitespace'): new = new.strip()
                    field = re.sub(r'[^a-z0-9]+', '_', headers[col-left].strip().lower()).strip('_')
                    if options.get('emails') and 'email' in field.replace('_', ''): new = new.lower().strip()
                    if options.get('phones') and any(t in field for t in ('phone', 'mobile', 'contact_number')): new = phone_cleaner(new)
                    if new != value:
                        changes[ref] = new
                        log.append({'Cell': ref, 'Before': value, 'After': new})
                elif value != '':
                    skipped += 1
                before.append(value); after.append(new)
                if str(new).strip() == '': missing.append({'Cell': ref, 'Column': labels[col-left]})
            key = tuple(after)
            if any(str(x).strip() for x in after):
                if key in seen: duplicates.append(row)
                seen.add(key)
            before_rows.append(before); rows.append(after)
        index = pd.Index(range(top+1, bottom+1), name='Excel row')
        reports[name] = {'original': pd.DataFrame(before_rows, columns=labels, index=index),
                         'cleaned': pd.DataFrame(rows, columns=labels, index=index),
                         'log': pd.DataFrame(log, columns=['Cell','Before','After']),
                         'missing': pd.DataFrame(missing, columns=['Cell','Column']),
                         'duplicates': pd.DataFrame({'Excel row': duplicates}),
                         'changed': len(changes), 'skipped': skipped, 'protected': sheet['protected']}
        if changes:
            patched_parts[source.sheets[name]] = _patch(sheet['raw'], changes)
    if not patched_parts:
        return source.raw, reports
    output = io.BytesIO()
    with zipfile.ZipFile(output, 'w') as z:
        z.comment = source.z.comment
        for item in source.z.infolist():
            z.writestr(item, patched_parts.get(item.filename, source.z.read(item.filename)))
    return output.getvalue(), reports
