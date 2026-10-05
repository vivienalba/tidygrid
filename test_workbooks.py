"""Regression tests use a small, literal OOXML fixture rather than a workbook writer."""
import io
import unittest
import zipfile
from xml.etree import ElementTree as ET

from openpyxl import load_workbook
from workbook_cleaner import WorkbookSource, clean_workbook, NS


def fixture(protected=False):
    parts = {
        '[Content_Types].xml': '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"><Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/><Default Extension="xml" ContentType="application/xml"/><Override PartName="/xl/workbook.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"/><Override PartName="/xl/worksheets/sheet1.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/><Override PartName="/xl/worksheets/sheet2.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/><Override PartName="/xl/styles.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.styles+xml"/><Override PartName="/xl/sharedStrings.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sharedStrings+xml"/></Types>',
        '_rels/.rels': '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="xl/workbook.xml"/></Relationships>',
        'xl/workbook.xml': f'<workbook xmlns="{NS}" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"><bookViews><workbookView/></bookViews><sheets><sheet name="Customers" sheetId="1" r:id="rId1"/><sheet name="Archive" sheetId="2" state="hidden" r:id="rId2"/></sheets><definedNames><definedName name="ContactCells">Customers!$A$2:$B$3</definedName></definedNames></workbook>',
        'xl/_rels/workbook.xml.rels': '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet" Target="worksheets/sheet1.xml"/><Relationship Id="rId2" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet" Target="worksheets/sheet2.xml"/><Relationship Id="rId3" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/styles" Target="styles.xml"/><Relationship Id="rId4" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/sharedStrings" Target="sharedStrings.xml"/></Relationships>',
        'xl/styles.xml': f'<styleSheet xmlns="{NS}"><numFmts count="1"><numFmt numFmtId="164" formatCode="00000"/></numFmts><fonts count="2"><font><name val="Calibri"/><sz val="11"/></font><font><b/><color rgb="FF6D3485"/><name val="Georgia"/><sz val="14"/></font></fonts><fills count="3"><fill><patternFill patternType="none"/></fill><fill><patternFill patternType="gray125"/></fill><fill><patternFill patternType="solid"><fgColor rgb="FFFFFFB5"/><bgColor indexed="64"/></patternFill></fill></fills><borders count="2"><border/><border><left style="thin"><color rgb="FF171219"/></left></border></borders><cellStyleXfs count="1"><xf numFmtId="0" fontId="0" fillId="0" borderId="0"/></cellStyleXfs><cellXfs count="3"><xf numFmtId="0" fontId="0" fillId="0" borderId="0" xfId="0"/><xf numFmtId="0" fontId="1" fillId="2" borderId="1" xfId="0"><alignment horizontal="center" wrapText="1"/></xf><xf numFmtId="164" fontId="0" fillId="0" borderId="0" xfId="0"/></cellXfs><cellStyles count="1"><cellStyle name="Normal" xfId="0" builtinId="0"/></cellStyles></styleSheet>',
        'xl/sharedStrings.xml': f'<sst xmlns="{NS}" count="3" uniqueCount="2"><si><t xml:space="preserve"> MARIA@EMAIL.COM </t></si><si><r><rPr><b/></rPr><t> RICH@EMAIL.COM </t></r></si></sst>',
        'xl/worksheets/sheet1.xml': f'''<worksheet xmlns="{NS}"><dimension ref="A1:F6"/><sheetViews><sheetView workbookViewId="0"><pane xSplit="1" ySplit="1" topLeftCell="B2" activePane="bottomRight" state="frozen"/></sheetView></sheetViews><cols><col min="1" max="1" width="28" customWidth="1"/></cols><sheetData>
        <row r="1" ht="32" customHeight="1"><c r="A1" s="1" t="inlineStr"><is><t>Email Address</t></is></c><c r="B1" t="inlineStr"><is><t>Mobile Number</t></is></c><c r="C1" t="inlineStr"><is><t>Code</t></is></c><c r="D1" t="inlineStr"><is><t>Calc</t></is></c></row>
        <row r="2" ht="24" customHeight="1"><c r="A2" s="1" t="s"><v>0</v></c><c r="B2" s="1" t="inlineStr"><is><t>0917-123-4567</t></is></c><c r="C2" s="2"><v>12</v></c><c r="D2"><f>C2*2</f><v>24</v></c></row>
        <row r="3"><c r="A3" t="inlineStr"><is><t>  MERGED@EMAIL.COM </t></is></c><c r="B3" t="inlineStr"><is><t> LINK@EMAIL.COM </t></is></c><c r="C3" t="s"><v>1</v></c></row>
        <row r="4"><c r="A4" t="inlineStr"><is><t> SAME@EMAIL.COM </t></is></c><c r="B4" t="inlineStr"><is><t>0918-111-2222</t></is></c></row>
        <row r="5"><c r="A5" t="inlineStr"><is><t>same@email.com</t></is></c><c r="B5" t="inlineStr"><is><t>+639181112222</t></is></c></row>
        <row r="6"><c r="A6" t="inlineStr"><is><t> OUTSIDE@EMAIL.COM </t></is></c></row></sheetData>{'<sheetProtection sheet="1"/>' if protected else ''}<autoFilter ref="A1:D5"/><mergeCells count="1"><mergeCell ref="A3:A3"/></mergeCells><conditionalFormatting sqref="C2:C5"><cfRule type="expression" priority="1"><formula>C2&gt;5</formula></cfRule></conditionalFormatting><dataValidations count="1"><dataValidation type="whole" sqref="C2:C5"><formula1>0</formula1><formula2>1000</formula2></dataValidation></dataValidations><hyperlinks><hyperlink ref="B3" location="Archive!A1"/></hyperlinks><pageMargins left="0.7" right="0.7" top="0.75" bottom="0.75" header="0.3" footer="0.3"/><pageSetup orientation="landscape" paperSize="9"/></worksheet>''',
        'xl/worksheets/sheet2.xml': f'<worksheet xmlns="{NS}"><sheetData><row r="1"><c r="A1" t="inlineStr"><is><t>Email Address</t></is></c></row><row r="2"><c r="A2" t="s"><v>0</v></c></row></sheetData></worksheet>',
        # Opaque unsupported feature parts must survive byte for byte too.
        'customXml/design-metadata.xml': '<design>unchanged opaque metadata</design>',
    }
    out = io.BytesIO()
    with zipfile.ZipFile(out, 'w', zipfile.ZIP_DEFLATED) as z:
        for k, v in parts.items(): z.writestr(k, v)
    return out.getvalue()


def phone(v):
    digits = ''.join(c for c in v if c.isdigit())
    return '+63'+digits[1:] if len(digits)==11 and digits.startswith('09') else v


class WorkbookTests(unittest.TestCase):
    def test_only_selected_plain_text_values_change(self):
        raw = fixture()
        out, reports = clean_workbook(WorkbookSource(raw), {'Customers':'A1:D5'}, {'whitespace':True,'emails':True,'phones':True}, phone)
        with zipfile.ZipFile(io.BytesIO(raw)) as before, zipfile.ZipFile(io.BytesIO(out)) as after:
            self.assertEqual(before.namelist(), after.namelist())
            for name in before.namelist():
                if name != 'xl/worksheets/sheet1.xml': self.assertEqual(before.read(name), after.read(name), name)
            a, b = ET.fromstring(before.read('xl/worksheets/sheet1.xml')), ET.fromstring(after.read('xl/worksheets/sheet1.xml'))
            for x,y in zip(a,b):
                if x.tag != f'{{{NS}}}sheetData': self.assertEqual(ET.tostring(x), ET.tostring(y))
            old_cells = {c.get('r'):c for c in a.iter(f'{{{NS}}}c')}
            new_cells = {c.get('r'):c for c in b.iter(f'{{{NS}}}c')}
            self.assertEqual(old_cells.keys(),new_cells.keys())
            for ref, c in old_cells.items():
                self.assertEqual(c.get('s'),new_cells[ref].get('s'), ref)
            for ref in ['A1','A3','B3','C3','C2','D2','A6']:
                self.assertEqual(ET.tostring(old_cells[ref]),ET.tostring(new_cells[ref]), ref)
        wb = load_workbook(io.BytesIO(out))
        self.assertEqual(wb['Customers']['A2'].value, 'maria@email.com')
        self.assertEqual(wb['Customers']['B2'].value, '+639171234567')
        self.assertEqual(wb['Customers']['C2'].number_format, '00000')
        self.assertEqual(wb['Customers']['D2'].value, '=C2*2')
        self.assertEqual(wb['Customers']['A2'].fill.fgColor.rgb, 'FFFFFFB5')
        self.assertEqual(wb['Customers'].freeze_panes,'B2')
        self.assertEqual(wb['Archive']['A2'].value, ' MARIA@EMAIL.COM ')
        self.assertEqual(wb['Archive'].sheet_state, 'hidden')
        self.assertEqual(reports['Customers']['duplicates']['Excel row'].tolist(), [5])
        self.assertEqual(len(reports['Customers']['cleaned']),4)

    def test_all_sheets_independent_and_shared_strings_untouched(self):
        raw=fixture()
        out,r=clean_workbook(WorkbookSource(raw), {'Customers':'A1:D5','Archive':'A1:A2'}, {'emails':True}, phone)
        wb=load_workbook(io.BytesIO(out))
        self.assertEqual(wb['Archive']['A2'].value,'maria@email.com')
        self.assertEqual(set(r),{'Customers','Archive'})

    def test_protected_sheet_and_noop_return_original_bytes(self):
        raw=fixture(True)
        out,r=clean_workbook(WorkbookSource(raw),{'Customers':'A1:D5'},{'whitespace':True,'emails':True,'phones':True},phone)
        self.assertEqual(raw,out)
        self.assertEqual(r['Customers']['changed'],0)
        raw=fixture()
        out,_=clean_workbook(WorkbookSource(raw),{'Customers':'A1:D5'},{},phone)
        self.assertEqual(raw,out)

    def test_invalid_ranges_and_signed_files_are_rejected(self):
        for region in ['A:A','A0:D5','D5:A1','A1:XFD1048576']:
            with self.assertRaises(ValueError): clean_workbook(WorkbookSource(fixture()),{'Customers':region},{},phone)
        out=io.BytesIO(fixture())
        with zipfile.ZipFile(out,'a') as z: z.writestr('_xmlsignatures/sig1.xml','signature')
        with self.assertRaises(ValueError): WorkbookSource(out.getvalue())

    def test_changed_text_is_safe_literal_not_formula(self):
        raw=fixture()
        # Replace the fixture through the package, preserving valid ZIP checksums.
        buf=io.BytesIO()
        with zipfile.ZipFile(io.BytesIO(raw)) as z, zipfile.ZipFile(buf,'w') as dest:
            for i in z.infolist():
                b=z.read(i.filename)
                if i.filename=='xl/sharedStrings.xml': b=b.replace(b' MARIA@EMAIL.COM ',b' =SUM(1,2)&amp;test ')
                dest.writestr(i,b)
        out,_=clean_workbook(WorkbookSource(buf.getvalue()),{'Archive':'A1:A2'},{'whitespace':True},phone)
        c=load_workbook(io.BytesIO(out))['Archive']['A2']
        self.assertEqual(c.value,'=SUM(1,2)&test')
        self.assertEqual(c.data_type,'s')


if __name__=='__main__': unittest.main()
